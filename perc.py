import openpyxl, docx, pyautogui, pyperclip, os
os.chdir('C:\\Users\\dario\\Desktop\\aika')
wb = openpyxl.Workbook()
sheet = wb.active
sheet["A1"]="a1"
sheet["A2"]="a2"
sheet["A3"]="a3"
sheet["B1"]="b1"
sheet["B2"]="b2"
sheet['B3']=1
sheet['B4']=2
sheet['C1']=sheet['B3'].value + sheet['B4'].value
sheet['c2']="lo de arriba es la suma de B3+B4"

pyautogui.moveTo(30,50,1)
print("El tamaño de la pantalla es: " + str(pyautogui.size()))
print("Estoy copiando en portapapeles esto: " + str(sheet["A1"].value) + " y esto: " + str(sheet["B3"].value))

d = docx.Document()
d.add_paragraph("Esto es un párrafo")
d.add_paragraph(sheet["A1"].value + " " + sheet["C2"].value + "." )
pyperclip.copy(sheet["A1"].value)
d.add_paragraph("Esto viene de pyperclip: " + pyperclip.paste())
pyperclip.copy(sheet["B3"].value)
d.add_paragraph("y esto también: " + pyperclip.paste())
d.save("elWord.docx")
wb.save("prueba.xlsx")

unDiccionario = {'llave1':'valor1','llave2':'valor2','key':'value'}
print(unDiccionario["key"])
print(list(unDiccionario.keys()))
print(list(unDiccionario.values()))
print(unDiccionario.items())
print(unDiccionario.get('keyoto', "bolas"))